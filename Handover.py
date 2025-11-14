from __future__ import annotations

import datetime as dt
import io
import json
import shutil
import sqlite3
import uuid
from collections import defaultdict, deque
from pathlib import Path
from urllib.parse import quote
from typing import Any, Dict, Iterable, List, Optional, Set, Tuple

import pandas as pd
import streamlit as st

try:  # Plotly is optional; degrade gracefully if unavailable
    import plotly.express as px  # type: ignore
except ImportError:  # pragma: no cover - visual enhancement only
    px = None  # type: ignore[assignment]


def _trigger_rerun() -> None:
    """Trigger a Streamlit rerun, handling API differences across versions."""
    if hasattr(st, "rerun"):
        st.rerun()
    elif hasattr(st, "experimental_rerun"):
        st.experimental_rerun()

try:  # Prefer xlsxwriter for richer formatting, otherwise fall back
    import xlsxwriter  # type: ignore  # noqa: F401

    EXCEL_ENGINE = "xlsxwriter"
except ImportError:  # pragma: no cover - platform dependent
    try:
        import openpyxl  # type: ignore  # noqa: F401

        EXCEL_ENGINE = "openpyxl"
    except ImportError:  # pragma: no cover - platform dependent
        EXCEL_ENGINE = None


DB_PATH = Path(__file__).with_name("handover_tracker.db")
BACKUP_DIR = DB_PATH.with_name("handover_backups")
STATUS_OPTIONS = ["Not Started", "In Progress", "Waiting", "Completed", "Archived"]
STATUS_BADGES = {
    "Not Started": "🟥 Not Started",
    "In Progress": "🟨 In Progress",
    "Waiting": "🟦 Waiting",
    "Completed": "🟩 Completed",
    "Archived": "⬜ Archived",
}

STAKEHOLDER_SUPPORT_LEVELS = [
    "Advocate",
    "Neutral",
    "Resistant",
]

STAKEHOLDER_SCALE = [1, 2, 3, 4, 5]

STATUS_COLORS = {
    "Not Started": "#d62728",
    "In Progress": "#ffbf00",
    "Waiting": "#1f77b4",
    "Completed": "#2ca02c",
    "Archived": "#7f7f7f",
}

ORG_LEVEL_COLORS = [
    "#FEE4E2",  # Level 0
    "#FEF3C7",  # Level 1
    "#DCFCE7",  # Level 2
    "#BFDBFE",  # Level 3
    "#E9D5FF",  # Level 4
    "#FBCFE8",  # Level 5+
]

try:
    from barfi.flow import Block
    from barfi.flow.streamlit import st_flow

    BARFI_AVAILABLE = True
except ImportError:  # pragma: no cover - optional dependency
    Block = None  # type: ignore[assignment]
    st_flow = None  # type: ignore[assignment]
    BARFI_AVAILABLE = False


def _graphviz_label(text: str) -> str:
    return text.replace("\\", "\\\\").replace("\n", "\\n").replace("\"", "\\\"")


def _wrap_title(title: str, max_chars: int = 40) -> str:
    clean = title.strip()
    if len(clean) <= max_chars:
        return clean

    tokens = clean.split()
    if not tokens:
        return clean

    lines: List[str] = []
    current: List[str] = []
    current_len = 0

    for token in tokens:
        next_len = current_len + len(token) + (1 if current else 0)
        if current and next_len > max_chars:
            lines.append(" ".join(current))
            current = [token]
            current_len = len(token)
        else:
            current.append(token)
            current_len = next_len

    if current:
        lines.append(" ".join(current))

    return "\n".join(lines)


def _org_level_color(level: int) -> str:
    if not ORG_LEVEL_COLORS:
        return "#EDF2F7"
    return ORG_LEVEL_COLORS[min(level, len(ORG_LEVEL_COLORS) - 1)]


def _estimate_node_width(text: str, min_width: float = 3.0, char_width: float = 0.22) -> float:
    lines = text.split("\n")
    longest = max((len(line) for line in lines), default=len(text))
    return max(min_width, longest * char_width)


def _barfi_base_blocks() -> List["Block"]:
    if not BARFI_AVAILABLE:
        return []

    step_block = Block(name="Step")
    step_block.add_input(name="In")
    step_block.add_output(name="Next")

    decision_block = Block(name="Decision")
    decision_block.add_input(name="In")
    decision_block.add_output(name="Yes")
    decision_block.add_output(name="No")

    end_block = Block(name="End")
    end_block.add_input(name="In")

    return [step_block, decision_block, end_block]


def render_barfi_editor(workflow_label: str, workflow_id: int) -> None:
    if not BARFI_AVAILABLE or st_flow is None:
        st.info(
            "Install `barfi[streamlit]` to unlock the drag-and-drop editor.\n"
            "Run: `pip install \"barfi[streamlit]\"`.",
            icon="ℹ️",
        )
        return

    st.caption(
        "Beta: Build or rearrange workflow steps visually. Drag blocks from the right-click menu, "
        "connect outputs to inputs, then review the generated schema below. Saving back to the register "
        "will be added in a follow-up iteration."
    )

    base_blocks = _barfi_base_blocks()
    if not base_blocks:
        st.warning("No Barfi blocks available.")
        return

    barfi_key = f"workflow_barfi_{workflow_id}"
    editor_result = st_flow(base_blocks, key=barfi_key)

    schema = getattr(editor_result, "editor_schema", None)
    if schema is None:
        st.info("Create nodes and click Execute in the editor to produce a schema.")
        return

    try:
        schema_payload = schema.dict()  # type: ignore[attr-defined]
    except AttributeError:
        try:
            schema_payload = schema.to_dict()  # type: ignore[attr-defined]
        except AttributeError:
            schema_payload = schema

    st.markdown("##### Generated Barfi schema")
    st.code(json.dumps(schema_payload, indent=2, default=str), language="json")
    st.info(
        "Schema persistence back to the workflow database is not yet wired. \n"
        "Use the above JSON as a reference or export template for now.",
        icon="🛠️",
    )


def build_graphviz_workflow(
    steps_df: pd.DataFrame,
    font_size: int = 26,
    current_step_id: Optional[int] = None,
) -> Optional[str]:
    if steps_df.empty:
        return None

    ordered = steps_df.sort_values("step_order")
    node_font = max(12, font_size)
    lines: List[str] = [
        "digraph Workflow {",
        "    rankdir=LR;",
        '    graph [splines=ortho, nodesep=0.4, ranksep=0.75, margin="0.25,0.25"];',
        f'    node [shape=rectangle, style="rounded,filled", fontname="Helvetica", fontsize={node_font}, fillcolor="#F7FAFC", color="#4A5568", fontcolor="#1A202C", fixedsize=false];',
        '    edge [fontname="Helvetica", fontsize=22, color="#4A5568"];',
    ]

    lines.append(
        f'    start [label="Start", shape=circle, style="filled", fillcolor="#C1E1C1", color="#2F855A", fontcolor="#1A202C", fontsize={node_font}];'
    )
    lines.append(
        f'    finish [label="Finish", shape=doublecircle, style="filled", fillcolor="#C1E1C1", color="#2F855A", fontcolor="#1A202C", fontsize={node_font}];'
    )

    def node_id(step_id: int) -> str:
        return f"step_{int(step_id)}"

    first_step_id: Optional[int] = None

    row_chunks: Dict[int, List[str]] = {}
    node_count = 0
    existing_nodes: Dict[int, str] = {}
    order_to_node: Dict[int, str] = {}
    step_order_map: Dict[int, int] = {}

    for row in ordered.itertuples():
        if pd.isna(row.step_id):
            continue

        step_id = int(row.step_id)
        node_name = node_id(step_id)
        wrapped_title = _wrap_title(row.step_title.strip())
        title = f"Step {int(row.step_order)}: {wrapped_title}"
        label = _graphviz_label(title)
        is_decision = bool(getattr(row, "no_step_id", None) and not pd.isna(row.no_step_id))
        shape = "diamond" if is_decision else "rectangle"
        border_color = "#4A5568"
        highlight = current_step_id is not None and step_id == current_step_id
        node_attrs = [
            f"label=\"{label}\"",
            f"shape={shape}",
            f"color=\"{border_color}\"",
            "penwidth=2",
            f"fontsize={node_font}",
        ]
        if highlight:
            node_attrs.extend(
                [
                    "style=\"rounded,filled\"",
                    "fillcolor=\"#FFF3BF\"",
                    "color=\"#F59E0B\"",
                    "penwidth=3",
                ]
            )
        lines.append(f"    {node_name} [{', '.join(node_attrs)}];")

        step_order_val = int(row.step_order)
        existing_nodes[step_id] = node_name
        order_to_node[step_order_val] = node_name
        step_order_map[step_id] = step_order_val

        if first_step_id is None:
            first_step_id = step_id

        row_index = node_count // 4
        row_chunks.setdefault(row_index, []).append(node_name)
        node_count += 1

    if first_step_id is None:
        lines.append('    start -> finish;')
    else:
        lines.append(f'    start -> {node_id(first_step_id)};')

    if row_chunks:
        for row_index, nodes_in_row in sorted(row_chunks.items()):
            if not nodes_in_row:
                continue

            node_list = "; ".join(nodes_in_row)
            lines.append(f"    {{ rank=same; {node_list}; }}")

            if len(nodes_in_row) > 1:
                for left, right in zip(nodes_in_row, nodes_in_row[1:]):
                    lines.append(f"    {left} -> {right} [style=invis, weight=1];")

        lines.append("    { rank=source; start; }")
        lines.append("    { rank=sink; finish; }")

    for row in ordered.itertuples():
        if pd.isna(row.step_id):
            continue

        current = node_id(int(row.step_id))
        yes_target = getattr(row, "yes_step_id", None)
        no_target = getattr(row, "no_step_id", None)
        has_no_branch = bool(no_target is not None and not pd.isna(no_target))

        outgoing_edges: Set[Tuple[Optional[str], Optional[str]]] = set()

        def edge(target_raw: Optional[object], label_text: Optional[str]) -> None:
            target_name: Optional[str] = None
            target_key: Optional[int] = None
            target_order: Optional[int] = None
            if target_raw is not None and not pd.isna(target_raw):
                try:
                    target_key = int(target_raw)
                except (TypeError, ValueError):
                    target_key = None

            if target_key is not None:
                target_name = existing_nodes.get(target_key)
                if target_name is not None:
                    target_order = step_order_map.get(target_key)
                if target_name is None:
                    target_name = order_to_node.get(target_key)
                    if target_name is not None:
                        target_order = target_key

            signature = (target_name, label_text)
            if signature in outgoing_edges:
                return
            outgoing_edges.add(signature)

            attrs: List[str] = []
            if label_text:
                safe_label = _graphviz_label(label_text)
                attrs.append(f'label="{safe_label}"')
                if label_text.lower() == "yes":
                    attrs.append('color="#2F855A"')
                    attrs.append('fontcolor="#2F855A"')
                elif label_text.lower() == "no":
                    attrs.append('color="#C53030"')
                    attrs.append('fontcolor="#C53030"')
                    attrs.append('labelfloat=true')
                    attrs.append('style="dotted"')
            if target_order is not None and target_order < int(row.step_order):
                attrs.append('constraint=false')
            if target_name is None:
                attr_part = f" [ {' ,'.join(attrs)} ]".replace(" [  ]", "") if attrs else ""
                lines.append(f'    {current} -> finish{attr_part};')
                return

            attr_part = f" [ {' ,'.join(attrs)} ]".replace(" [  ]", "") if attrs else ""
            lines.append(f'    {current} -> {target_name}{attr_part};')

        if has_no_branch:
            edge(yes_target, "Yes")
            edge(no_target, "No")
        else:
            edge(yes_target, None)

    lines.append('}')
    return "\n".join(lines)


def status_badge(status: str) -> str:
    return STATUS_BADGES.get(status, status)


def status_color(status: str) -> str:
    return STATUS_COLORS.get(status, "#636EFA")


def ensure_workflow_tables(conn: sqlite3.Connection) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS workflow_processes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            owner TEXT,
            description TEXT,
            created_at TEXT NOT NULL
        )
        """
    )

    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS workflow_steps (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            workflow_id INTEGER NOT NULL,
            step_order INTEGER NOT NULL,
            title TEXT NOT NULL,
            status TEXT NOT NULL,
            yes_step_id INTEGER,
            no_step_id INTEGER,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (workflow_id) REFERENCES workflow_processes(id) ON DELETE CASCADE
        )
        """
    )

    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_workflow_steps_workflow
        ON workflow_steps (workflow_id, step_order)
        """
    )
    _ensure_workflow_step_branch_columns(conn)
    conn.commit()


def _ensure_workflow_step_branch_columns(conn: sqlite3.Connection) -> None:
    existing_columns = {
        row[1] for row in conn.execute("PRAGMA table_info('workflow_steps')")
    }
    if "yes_step_id" not in existing_columns:
        conn.execute("ALTER TABLE workflow_steps ADD COLUMN yes_step_id INTEGER")
    if "no_step_id" not in existing_columns:
        conn.execute("ALTER TABLE workflow_steps ADD COLUMN no_step_id INTEGER")


def ensure_topic_schema(conn: sqlite3.Connection) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS handover_topics (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            person TEXT NOT NULL,
            topic TEXT NOT NULL,
            meeting_date TEXT NOT NULL,
            details TEXT,
            status TEXT NOT NULL,
            next_steps TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )


def ensure_stakeholder_schema(conn: sqlite3.Connection) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS stakeholders (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            position TEXT,
            department TEXT,
            role_description TEXT,
            influence INTEGER NOT NULL,
            interest INTEGER NOT NULL,
            power INTEGER NOT NULL,
            support_level TEXT NOT NULL,
            notes TEXT,
            contact_id INTEGER,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (contact_id) REFERENCES contacts(id) ON DELETE SET NULL
        )
        """
    )


def _ensure_stakeholder_columns(conn: sqlite3.Connection) -> None:
    ensure_stakeholder_schema(conn)
    existing_columns = {
        row[1] for row in conn.execute("PRAGMA table_info('stakeholders')")
    }
    if "contact_id" not in existing_columns:
        conn.execute("ALTER TABLE stakeholders ADD COLUMN contact_id INTEGER")
        existing_columns = {
            row[1] for row in conn.execute("PRAGMA table_info('stakeholders')")
        }
    expected = {
        "name",
        "position",
        "department",
        "role_description",
        "influence",
        "interest",
        "power",
        "support_level",
        "notes",
        "contact_id",
        "created_at",
        "updated_at",
    }
    missing = expected - existing_columns
    if missing:
        raise RuntimeError(
            "Stakeholder table is missing expected columns: " + ", ".join(sorted(missing))
        )


def ensure_contact_schema(conn: sqlite3.Connection) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS contacts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            full_name TEXT NOT NULL,
            title TEXT,
            department TEXT,
            email TEXT,
            phone TEXT,
            location TEXT,
            manager_id INTEGER,
            notes TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (manager_id) REFERENCES contacts(id) ON DELETE SET NULL
        )
        """
    )


def _init_db_impl() -> None:
    """Ensure the handover and workflow tables exist."""

    with get_connection() as conn:
        ensure_topic_schema(conn)
        ensure_stakeholder_schema(conn)
        _ensure_stakeholder_columns(conn)
        ensure_contact_schema(conn)
        ensure_workflow_tables(conn)
        conn.commit()


def init_db() -> None:
    """Ensure core database structures are ready for use."""

    _init_db_impl()


def add_stakeholder(
    name: str,
    position: str,
    department: str,
    role_description: str,
    influence: int,
    interest: int,
    power: int,
    support_level: str,
    notes: str = "",
    contact_id: Optional[int] = None,
) -> None:
    now_iso = dt.datetime.utcnow().isoformat(timespec="seconds")
    payload = (
        name.strip(),
        position.strip(),
        department.strip(),
        role_description.strip(),
        int(influence),
        int(interest),
        int(power),
        support_level,
        notes.strip(),
        int(contact_id) if contact_id is not None else None,
        now_iso,
        now_iso,
    )

    with get_connection() as conn:
        _ensure_stakeholder_columns(conn)
        conn.execute(
            """
            INSERT INTO stakeholders (
                name, position, department, role_description, influence, interest, power,
                support_level, notes, contact_id, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            payload,
        )


def update_stakeholder(stakeholder_id: int, updates: Dict[str, Any]) -> None:
    if not updates:
        return

    set_parts: List[str] = []
    bindings: List[Any] = []

    for column, value in updates.items():
        set_parts.append(f"{column} = ?")
        bindings.append(value)

    set_parts.append("updated_at = ?")
    bindings.append(dt.datetime.utcnow().isoformat(timespec="seconds"))
    bindings.append(stakeholder_id)

    with get_connection() as conn:
        _ensure_stakeholder_columns(conn)
        conn.execute(
            f"UPDATE stakeholders SET {', '.join(set_parts)} WHERE id = ?",
            bindings,
        )


def delete_stakeholders(stakeholder_ids: Iterable[int]) -> None:
    ids = [int(sid) for sid in stakeholder_ids]
    if not ids:
        return

    placeholders = ",".join("?" for _ in ids)
    with get_connection() as conn:
        _ensure_stakeholder_columns(conn)
        conn.execute(
            f"DELETE FROM stakeholders WHERE id IN ({placeholders})",
            ids,
        )


def fetch_stakeholders() -> pd.DataFrame:
    with get_connection() as conn:
        _ensure_stakeholder_columns(conn)
        df = pd.read_sql_query(
            "SELECT * FROM stakeholders ORDER BY name COLLATE NOCASE",
            conn,
        )

    if not df.empty:
        df["created_at"] = pd.to_datetime(df["created_at"])
        df["updated_at"] = pd.to_datetime(df["updated_at"])
        if "contact_id" in df.columns:
            df["contact_id"] = pd.to_numeric(df["contact_id"], errors="coerce").astype("Int64")

    return df


@st.cache_data(show_spinner=False)
def load_stakeholders() -> pd.DataFrame:
    return fetch_stakeholders()


def clear_cached_stakeholders() -> None:
    load_stakeholders.clear()  # type: ignore[attr-defined]


def add_contact(
    full_name: str,
    title: str,
    department: str,
    email: str,
    phone: str,
    location: str,
    manager_id: Optional[int],
    notes: str,
) -> None:
    now_iso = dt.datetime.utcnow().isoformat(timespec="seconds")
    payload: Tuple[object, ...] = (
        full_name.strip(),
        title.strip(),
        department.strip(),
        email.strip(),
        phone.strip(),
        location.strip(),
        int(manager_id) if manager_id is not None else None,
        notes.strip(),
        now_iso,
        now_iso,
    )

    with get_connection() as conn:
        ensure_contact_schema(conn)
        conn.execute(
            """
            INSERT INTO contacts (
                full_name, title, department, email, phone, location, manager_id, notes, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            payload,
        )


def update_contact(contact_id: int, updates: Dict[str, Any]) -> None:
    if not updates:
        return

    set_parts: List[str] = []
    bindings: List[Any] = []

    for column, value in updates.items():
        set_parts.append(f"{column} = ?")
        bindings.append(value)

    set_parts.append("updated_at = ?")
    bindings.append(dt.datetime.utcnow().isoformat(timespec="seconds"))
    bindings.append(contact_id)

    with get_connection() as conn:
        ensure_contact_schema(conn)
        conn.execute(
            f"UPDATE contacts SET {', '.join(set_parts)} WHERE id = ?",
            bindings,
        )


def delete_contacts(contact_ids: Iterable[int]) -> None:
    ids = [int(cid) for cid in contact_ids]
    if not ids:
        return

    placeholders = ",".join("?" for _ in ids)
    with get_connection() as conn:
        ensure_contact_schema(conn)
        conn.execute(
            f"DELETE FROM contacts WHERE id IN ({placeholders})",
            ids,
        )


def fetch_contacts() -> pd.DataFrame:
    with get_connection() as conn:
        ensure_contact_schema(conn)
        df = pd.read_sql_query(
            "SELECT * FROM contacts ORDER BY full_name COLLATE NOCASE",
            conn,
        )

    if not df.empty:
        df["created_at"] = pd.to_datetime(df["created_at"])
        df["updated_at"] = pd.to_datetime(df["updated_at"])

    return df


@st.cache_data(show_spinner=False)
def load_contacts() -> pd.DataFrame:
    return fetch_contacts()


def clear_cached_contacts() -> None:
    load_contacts.clear()  # type: ignore[attr-defined]


def build_excel_report(df: pd.DataFrame) -> bytes:
    """Generate an Excel workbook with handover data."""

    export_df = df.copy()
    export_df["Meeting Date"] = export_df["meeting_date"].apply(lambda d: d.strftime("%d %b %Y"))
    export_df["Status"] = export_df["status"].apply(status_badge)
    export_df["Created"] = export_df["created_at"].dt.strftime("%d %b %Y %H:%M")
    export_df["Last Updated"] = export_df["updated_at"].dt.strftime("%d %b %Y %H:%M")

    columns = [
        "person",
        "topic",
        "Meeting Date",
        "details",
        "Status",
        "next_steps",
        "Created",
        "Last Updated",
    ]

    export_df = export_df[columns]
    export_df = export_df.rename(
        columns={
            "person": "Owner",
            "topic": "Topic",
            "details": "Key Details",
            "next_steps": "Next Steps",
        }
    )

    buffer = io.BytesIO()
    if EXCEL_ENGINE is None:
        raise RuntimeError("No Excel writer engine available. Install xlsxwriter or openpyxl.")

    with pd.ExcelWriter(buffer, engine=EXCEL_ENGINE) as writer:
        export_df.to_excel(writer, index=False, sheet_name="Handover Register")
        if EXCEL_ENGINE == "xlsxwriter":
            workbook = writer.book
            worksheet = writer.sheets["Handover Register"]

            header_format = workbook.add_format(
                {"bold": True, "bg_color": "#003366", "font_color": "#FFFFFF", "align": "center"}
            )
            for col_num, value in enumerate(export_df.columns):
                worksheet.write(0, col_num, value, header_format)
                column_width = max(15, int(export_df[value].astype(str).str.len().max()) + 2)
                worksheet.set_column(col_num, col_num, min(column_width, 60))

            date_format = workbook.add_format({"num_format": "dd mmm yyyy"})
            worksheet.set_column("C:C", 18, date_format)

    buffer.seek(0)
    return buffer.read()


def build_summary(df: pd.DataFrame) -> str:
    total = len(df)
    completed = int((df["status"] == "Completed").sum())
    waiting = int((df["status"] == "Waiting").sum())
    in_progress = int((df["status"] == "In Progress").sum())

    upcoming_window = dt.date.today() + dt.timedelta(days=7)
    upcoming = int((df["meeting_date"] <= upcoming_window).sum())

    summary_lines = [
        "# Handover summary",
        f"- Total topics: {len(df)}",
        f"- Completed: {completed}",
        f"- In progress: {in_progress}",
        f"- Waiting on others: {waiting}",
        f"- Due within 7 days: {upcoming}",
        "",
        "## Highlights",
    ]

    top_items = (
        df.sort_values("meeting_date")
        .head(5)
        .assign(
            badge=lambda d: d["status"].apply(status_badge),
            due=lambda d: d["meeting_date"].apply(lambda date: date.strftime("%d %b")),
        )
    )

    if top_items.empty:
        summary_lines.append("No topics recorded yet.")
    else:
        for row in top_items.itertuples():
            summary_lines.append(
                f"- {row.due}: {row.topic} ({row.person}) · {row.badge}\n  Next: {row.next_steps or 'No next steps logged.'}"
            )

    summary_lines.append("")
    summary_lines.append("Generated via the Handover Tracker")
    return "\n".join(summary_lines)


def get_connection() -> sqlite3.Connection:
    """Return a SQLite connection, creating the database if needed."""

    DB_PATH.touch(exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db() -> None:
    """Ensure the handover table exists."""

    with get_connection() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS handover_topics (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                person TEXT NOT NULL,
                topic TEXT NOT NULL,
                meeting_date TEXT NOT NULL,
                details TEXT,
                status TEXT NOT NULL,
                next_steps TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            )
            """
        )


def add_topic(
    person: str,
    topic: str,
    meeting_date: dt.date,
    details: str,
    status: str,
    next_steps: str,
) -> None:
    now_iso = dt.datetime.utcnow().isoformat(timespec="seconds")
    payload = (
        person.strip(),
        topic.strip(),
        meeting_date.isoformat(),
        details.strip(),
        status,
        next_steps.strip(),
        now_iso,
        now_iso,
    )

    with get_connection() as conn:
        ensure_topic_schema(conn)
        conn.execute(
            """
            INSERT INTO handover_topics (
                person, topic, meeting_date, details, status, next_steps, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            payload,
        )


def update_topic(
    topic_id: int,
    updates: Dict[str, Any],
    meeting_date: Optional[dt.date] = None,
) -> None:
    set_parts: list[str] = []
    bindings: List[Any] = []

    for column, value in updates.items():
        set_parts.append(f"{column} = ?")
        bindings.append(value)

    if meeting_date is not None:
        set_parts.append("meeting_date = ?")
        bindings.append(meeting_date.isoformat())

    set_parts.append("updated_at = ?")
    bindings.append(dt.datetime.utcnow().isoformat(timespec="seconds"))
    bindings.append(topic_id)

    with get_connection() as conn:
        ensure_topic_schema(conn)
        conn.execute(
            f"UPDATE handover_topics SET {' ,'.join(set_parts)} WHERE id = ?",
            bindings,
        )


def fetch_topics(filters: Optional[Dict[str, Iterable[str]]] = None) -> pd.DataFrame:
    query = (
        "SELECT id, person, topic, meeting_date, details, status, next_steps, created_at, updated_at "
        "FROM handover_topics"
    )
    conditions: list[str] = []
    params: list[str] = []

    if filters:
        if persons := list(filters.get("person", [])):
            placeholders = ",".join("?" for _ in persons)
            conditions.append(f"person IN ({placeholders})")
            params.extend(persons)

        if statuses := list(filters.get("status", [])):
            placeholders = ",".join("?" for _ in statuses)
            conditions.append(f"status IN ({placeholders})")
            params.extend(statuses)

        if conditions:
            query += " WHERE " + " AND ".join(conditions)

    query += " ORDER BY meeting_date ASC"

    with get_connection() as conn:
        ensure_topic_schema(conn)
        df = pd.read_sql_query(query, conn, params=params)

    if not df.empty:
        df["meeting_date"] = pd.to_datetime(df["meeting_date"]).dt.date
        df["created_at"] = pd.to_datetime(df["created_at"])
        df["updated_at"] = pd.to_datetime(df["updated_at"])

    return df


def fetch_topic(topic_id: int) -> Optional[pd.Series]:
    df = fetch_topics()
    if df.empty:
        return None
    match = df.loc[df["id"] == topic_id]
    if match.empty:
        return None
    return match.iloc[0]


@st.cache_data(show_spinner=False)
def load_topics(filters: Dict[str, Iterable[str]]) -> pd.DataFrame:
    return fetch_topics(filters)


def clear_cached_topics() -> None:
    load_topics.clear()  # type: ignore[attr-defined]


def add_workflow_process(name: str, owner: str, description: str, steps: List[str]) -> None:
    init_db()
    now_iso = dt.datetime.utcnow().isoformat(timespec="seconds")
    payload = (name.strip(), owner.strip(), description.strip(), now_iso)

    for attempt in (1, 2):
        try:
            with get_connection() as conn:
                ensure_workflow_tables(conn)
                cursor = conn.execute(
                    """
                    INSERT INTO workflow_processes (name, owner, description, created_at)
                    VALUES (?, ?, ?, ?)
                    """,
                    payload,
                )
                workflow_id = cursor.lastrowid

                inserted_ids: List[int] = []
                for order, title in enumerate(steps, start=1):
                    step_cursor = conn.execute(
                        """
                        INSERT INTO workflow_steps (
                            workflow_id, step_order, title, status, yes_step_id, no_step_id, created_at, updated_at
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            workflow_id,
                            order,
                            title.strip(),
                            "Not Started",
                            None,
                            None,
                            now_iso,
                            now_iso,
                        ),
                    )
                    inserted_ids.append(int(step_cursor.lastrowid))

                for idx, step_id in enumerate(inserted_ids):
                    yes_target = inserted_ids[idx + 1] if idx + 1 < len(inserted_ids) else None
                    conn.execute(
                        "UPDATE workflow_steps SET yes_step_id = ? WHERE id = ?",
                        (yes_target, step_id),
                    )
            break
        except sqlite3.OperationalError as exc:
            if attempt == 2 or "workflow_processes" not in str(exc):
                raise
            init_db()


def fetch_workflows() -> pd.DataFrame:
    init_db()
    query = """
        SELECT
            w.id AS workflow_id,
            w.name AS workflow_name,
            w.owner AS workflow_owner,
            w.description AS workflow_description,
            w.created_at AS workflow_created_at,
            s.id AS step_id,
            s.step_order,
            s.title AS step_title,
            s.status AS step_status,
            s.yes_step_id,
            s.no_step_id,
            s.updated_at AS step_updated_at
        FROM workflow_processes w
        LEFT JOIN workflow_steps s ON s.workflow_id = w.id
        ORDER BY w.created_at DESC, s.step_order ASC
    """

    try:
        with get_connection() as conn:
            df = pd.read_sql_query(query, conn)
    except (sqlite3.OperationalError, pd.errors.DatabaseError) as exc:
        if "workflow_processes" in str(exc):
            return pd.DataFrame(
                columns=[
                    "workflow_id",
                    "workflow_name",
                    "workflow_owner",
                    "workflow_description",
                    "workflow_created_at",
                    "step_id",
                    "step_order",
                    "step_title",
                    "step_status",
                    "yes_step_id",
                    "no_step_id",
                    "step_updated_at",
                ]
            )
        raise

    if not df.empty:
        df["workflow_created_at"] = pd.to_datetime(df["workflow_created_at"])
        if "step_updated_at" in df.columns:
            df["step_updated_at"] = pd.to_datetime(df["step_updated_at"])

    return df


@st.cache_data(show_spinner=False)
def load_workflows() -> pd.DataFrame:
    return fetch_workflows()


def clear_cached_workflows() -> None:
    load_workflows.clear()  # type: ignore[attr-defined]


def ensure_backup_directory() -> Path:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    return BACKUP_DIR


def _backup_filename(timestamp: dt.datetime) -> str:
    return f"handover_backup_{timestamp.strftime('%Y%m%d_%H%M%S')}.db"


def _parse_backup_timestamp(path: Path) -> Optional[dt.datetime]:
    stem = path.stem
    prefix = "handover_backup_"
    if not stem.startswith(prefix):
        return None
    raw = stem[len(prefix) :]
    try:
        return dt.datetime.strptime(raw, "%Y%m%d_%H%M%S")
    except ValueError:
        return None


def list_database_backups() -> List[Path]:
    if not BACKUP_DIR.exists():
        return []
    return sorted(BACKUP_DIR.glob("handover_backup_*.db"), reverse=True)


def create_database_backup() -> Path:
    ensure_backup_directory()
    init_db()
    timestamp = dt.datetime.utcnow()
    backup_path = BACKUP_DIR / _backup_filename(timestamp)
    shutil.copy2(DB_PATH, backup_path)
    return backup_path


def restore_database_backup(backup_path: Path) -> None:
    if not backup_path.exists() or not backup_path.is_file():
        raise FileNotFoundError(f"Backup not found: {backup_path}")

    shutil.copy2(backup_path, DB_PATH)
    init_db()
    clear_cached_topics()
    clear_cached_contacts()
    clear_cached_stakeholders()
    clear_cached_workflows()


def trigger_rerun() -> None:
    rerun = getattr(st, "rerun", None)
    if callable(rerun):
        rerun()
        return

    experimental_rerun = getattr(st, "experimental_rerun", None)
    if callable(experimental_rerun):
        experimental_rerun()


def update_workflow_steps(step_updates: Dict[int, Dict[str, object]]) -> None:
    if not step_updates:
        return

    init_db()
    now_iso = dt.datetime.utcnow().isoformat(timespec="seconds")
    for attempt in (1, 2):
        try:
            with get_connection() as conn:
                ensure_workflow_tables(conn)
                for step_id, fields in step_updates.items():
                    if not fields:
                        continue

                    columns: List[str] = []
                    values: List[object] = []

                    for column in ("title", "step_order", "yes_step_id", "no_step_id"):
                        if column in fields:
                            if column == "title" and isinstance(fields[column], str):
                                values.append(fields[column].strip())
                            else:
                                values.append(fields[column])
                            columns.append(f"{column} = ?")

                    if not columns:
                        continue

                    columns.append("updated_at = ?")
                    values.append(now_iso)
                    values.append(step_id)

                    conn.execute(
                        f"UPDATE workflow_steps SET {', '.join(columns)} WHERE id = ?",
                        values,
                    )
            break
        except sqlite3.OperationalError as exc:
            if attempt == 2 or "workflow_steps" not in str(exc):
                raise
            init_db()


def insert_workflow_step(
    workflow_id: int,
    step_order: int,
    title: str,
    yes_step_id: Optional[int],
    no_step_id: Optional[int],
) -> int:
    init_db()
    now_iso = dt.datetime.utcnow().isoformat(timespec="seconds")
    payload = (
        int(workflow_id),
        int(step_order),
        title.strip(),
        "Not Started",
        yes_step_id,
        no_step_id,
        now_iso,
        now_iso,
    )

    for attempt in (1, 2):
        try:
            with get_connection() as conn:
                ensure_workflow_tables(conn)
                cursor = conn.execute(
                    """
                    INSERT INTO workflow_steps (
                        workflow_id, step_order, title, status, yes_step_id, no_step_id, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    payload,
                )
                return int(cursor.lastrowid)
        except sqlite3.OperationalError as exc:
            if attempt == 2 or "workflow_steps" not in str(exc):
                raise
            init_db()
    return 0


def delete_workflow_steps(step_ids: Iterable[int]) -> None:
    ids = [int(step_id) for step_id in step_ids]
    if not ids:
        return

    init_db()
    for attempt in (1, 2):
        try:
            with get_connection() as conn:
                ensure_workflow_tables(conn)
                placeholders = ",".join("?" for _ in ids)
                conn.execute(
                    f"DELETE FROM workflow_steps WHERE id IN ({placeholders})",
                    ids,
                )
            break
        except sqlite3.OperationalError as exc:
            if attempt == 2 or "workflow_steps" not in str(exc):
                raise
            init_db()


def delete_workflow_process(workflow_id: int) -> None:
    init_db()
    for attempt in (1, 2):
        try:
            with get_connection() as conn:
                ensure_workflow_tables(conn)
                conn.execute(
                    "DELETE FROM workflow_processes WHERE id = ?",
                    (int(workflow_id),),
                )
            break
        except sqlite3.OperationalError as exc:
            if attempt == 2 or "workflow_processes" not in str(exc):
                raise
            init_db()


def render_workflow_diagram(steps_df: pd.DataFrame, font_size: int, current_step_id: Optional[int]) -> None:
    graphviz_code = build_graphviz_workflow(
        steps_df,
        font_size=font_size,
        current_step_id=current_step_id,
    )
    if graphviz_code:
        st.graphviz_chart(graphviz_code, use_container_width=True)


def render_workflows() -> None:
    st.subheader("🧭 Process workflow builder")
    st.markdown(
        "Establish repeatable handover steps and track their progress visually.")

    with st.form("create_workflow_form", clear_on_submit=True):
        name = st.text_input("Workflow name", placeholder="e.g., New manager onboarding")
        owner = st.text_input("Owner (optional)", placeholder="Who leads this process?")
        description = st.text_area(
            "Description",
            placeholder="What is the goal and scope of this workflow?",
        )
        steps_text = st.text_area(
            "Workflow steps (one per line)",
            placeholder="Kick-off alignment\nTool access handover\nShadowing and sign-off",
        )
        create_submitted = st.form_submit_button("Create workflow", type="primary")

    if create_submitted:
        steps = [line.strip() for line in steps_text.splitlines() if line.strip()]
        if not name.strip():
            st.error("Workflow name is required.")
        elif not steps:
            st.error("Add at least one workflow step.")
        else:
            add_workflow_process(name, owner, description, steps)
            clear_cached_workflows()
            st.success("Workflow created.")
            trigger_rerun()

    df = load_workflows()

    if df.empty or df["workflow_id"].isna().all():
        st.info("No workflows captured yet. Use the form above to build your first one.")
        return

    with st.expander("Manage workflows", expanded=False):
        unique_workflows = (
            df.dropna(subset=["workflow_id"]).sort_values("workflow_created_at", ascending=False)
        )
        workflow_choices = {}
        for row in unique_workflows.itertuples():
            label_bits = [row.workflow_name]
            if isinstance(row.workflow_owner, str) and row.workflow_owner.strip():
                label_bits.append(f"Owner: {row.workflow_owner.strip()}")
            created_at = getattr(row, "workflow_created_at", None)
            if pd.notna(created_at):
                label_bits.append(created_at.strftime("%d %b %Y"))
            label = " · ".join(label_bits)
            workflow_choices[label] = int(row.workflow_id)

        if workflow_choices:
            col_pick, col_delete = st.columns([3, 1])
            with col_pick:
                selected_label = st.selectbox(
                    "Select a workflow to delete",
                    options=list(workflow_choices.keys()),
                    key="workflow_delete_select",
                )
            with col_delete:
                if st.button("Delete", key="workflow_delete_button"):
                    delete_workflow_process(workflow_choices[selected_label])
                    clear_cached_workflows()
                    st.success("Workflow removed.")
                    trigger_rerun()
        else:
            st.info("No workflows available to delete.")

    if "workflow_font_size" not in st.session_state:
        st.session_state["workflow_font_size"] = 26
    if "workflow_font_input" not in st.session_state:
        st.session_state["workflow_font_input"] = str(st.session_state["workflow_font_size"])
    if st.session_state.pop("workflow_font_reset_pending", False):
        st.session_state["workflow_font_size"] = 26
        st.session_state["workflow_font_input"] = "26"

    font_col, reset_col = st.columns([3, 1])
    with font_col:
        st.caption("Workflow diagram font size (12–60 pt)")
        font_input = st.text_input(
            "Font size (pt)",
            key="workflow_font_input",
            help="Type a number between 12 and 60 to adjust the workflow diagram font size.",
        )

        current_size = st.session_state["workflow_font_size"]
        font_input_str = font_input.strip()
        if font_input_str:
            try:
                parsed_size = int(float(font_input_str))
            except ValueError:
                st.warning("Enter a numeric font size, e.g., 24.")
            else:
                if 12 <= parsed_size <= 60:
                    if parsed_size != current_size:
                        st.session_state["workflow_font_size"] = parsed_size
                else:
                    st.warning("Font size must be between 12 and 60 pt.")

        st.write(f"**Current: {st.session_state['workflow_font_size']} pt**")

    with reset_col:
        if st.button("Reset", key="workflow_font_reset"):
            st.session_state["workflow_font_reset_pending"] = True
            _trigger_rerun()

    font_size = st.session_state["workflow_font_size"]

    grouped = df.groupby("workflow_id", sort=False)
    for workflow_id, group in grouped:
        workflow = group.iloc[0]
        steps_df = group.dropna(subset=["step_id"]).sort_values("step_order")

        header = workflow["workflow_name"]
        meta_bits: List[str] = []
        if isinstance(workflow.get("workflow_owner"), str) and workflow["workflow_owner"].strip():
            meta_bits.append(f"Owner: {workflow['workflow_owner'].strip()}")
        created_at = workflow.get("workflow_created_at")
        if pd.notna(created_at):
            meta_bits.append(f"Created {created_at.strftime('%d %b %Y')}")

        expander_label = header
        if meta_bits:
            expander_label = f"{header} · {' · '.join(meta_bits)}"

        with st.expander(expander_label, expanded=False):
            description = workflow.get("workflow_description")
            if isinstance(description, str) and description.strip():
                st.markdown(description.strip())

            tab_visual, tab_barfi = st.tabs([
                "Diagram & table",
                "Barfi editor (beta)",
            ])

            with tab_visual:
                state_key_current = f"workflow_current_step_{workflow_id}"
                state_key_source = f"workflow_current_source_{workflow_id}"
                if state_key_current not in st.session_state:
                    st.session_state[state_key_current] = None
                if state_key_source not in st.session_state:
                    st.session_state[state_key_source] = "diagram"

                current_step_id = st.session_state[state_key_current]

                ordered_steps = steps_df.sort_values("step_order") if not steps_df.empty else pd.DataFrame()
                step_options = (
                    ordered_steps["step_id"].dropna().astype(int).tolist()
                    if not ordered_steps.empty
                    else []
                )

                if step_options:
                    display_options = {
                        int(step_id): f"Step {int(step_order)}"
                        for step_id, step_order in zip(
                            ordered_steps["step_id"],
                            ordered_steps["step_order"],
                        )
                    }
                    selected_option = st.selectbox(
                        "Highlight step",
                        options=[None] + step_options,
                        index=(step_options.index(current_step_id) + 1) if current_step_id in step_options else 0,
                        format_func=lambda opt: "— None —" if opt is None else display_options.get(int(opt), str(opt)),
                        key=f"workflow_step_select_{workflow_id}",
                    )
                    st.session_state[state_key_current] = selected_option
                    st.session_state[state_key_source] = "selectbox"
                elif steps_df.empty:
                    st.caption("Add steps to enable highlighting.")

                if steps_df.empty:
                    st.info("No steps logged for this workflow yet. Use the table below to add steps.")
                else:
                    render_workflow_diagram(
                        steps_df,
                        font_size=font_size,
                        current_step_id=current_step_id,
                    )

                base_columns = ["step_id", "step_order", "step_title", "yes_step_id", "no_step_id"]
                if steps_df.empty:
                    table_source = pd.DataFrame(columns=base_columns)
                else:
                    table_source = steps_df[base_columns].copy()

                table_source = table_source.rename(
                    columns={
                        "step_id": "step_id",
                        "step_order": "step_order",
                        "step_title": "title",
                        "yes_step_id": "yes_step_id",
                        "no_step_id": "no_step_id",
                    }
                )

                if "step_id" in table_source.columns:
                    table_source["step_id"] = table_source["step_id"].astype("Int64")

                row_count = max(len(table_source), 1)
                table_height = min(700, 60 * row_count + 120)

                title_len = int(table_source["title"].astype(str).str.len().max()) if not table_source.empty else 0
                step_len = int(table_source["step_order"].astype(str).str.len().max()) if not table_source.empty else 1
                yes_len = int(table_source["yes_step_id"].astype(str).str.len().max()) if not table_source.empty else 1
                no_len = int(table_source["no_step_id"].astype(str).str.len().max()) if not table_source.empty else 1

                def _pick_width(length: int) -> str:
                    if length >= 25:
                        return "large"
                    if length >= 12:
                        return "medium"
                    return "small"

                editor_df = st.data_editor(
                    table_source,
                    num_rows="dynamic",
                    use_container_width=True,
                    height=table_height,
                    key=f"workflow_table_editor_{workflow_id}",
                    column_config={
                        "step_id": st.column_config.NumberColumn("Step ID", disabled=True, width="small"),
                        "step_order": st.column_config.NumberColumn(
                            "Step order", min_value=1, width=_pick_width(step_len)
                        ),
                        "title": st.column_config.TextColumn("Title", width=_pick_width(title_len)),
                        "yes_step_id": st.column_config.NumberColumn(
                            "Yes → Step ID", min_value=1, required=False, width=_pick_width(yes_len)
                        ),
                        "no_step_id": st.column_config.NumberColumn(
                            "No → Step ID", min_value=1, required=False, width=_pick_width(no_len)
                        ),
                    },
                    hide_index=True,
                )

                if st.button("Save table changes", key=f"workflow_table_save_{workflow_id}"):
                    if editor_df.empty:
                        st.error("Add at least one step before saving.")
                    else:
                        try:
                            valid_step_ids: Set[int] = set(
                                table_source["step_id"].dropna().astype(int).tolist()
                            )
                            original_ids = set(
                                table_source["step_id"].dropna().astype(int).tolist()
                            )
                            edited_ids = set(
                                editor_df["step_id"].dropna().astype(int).tolist()
                            )
                            to_delete = original_ids - edited_ids

                            updates: Dict[int, Dict[str, object]] = {}
                            inserts: List[Dict[str, object]] = []

                            for row in editor_df.itertuples(index=False):
                                step_id = row.step_id
                                title_raw = row.title
                                order_raw = row.step_order

                                if not isinstance(title_raw, str) or not title_raw.strip():
                                    raise ValueError("Every step must have a title.")
                                if pd.isna(order_raw):
                                    raise ValueError("Every step must have an order number.")

                                yes_raw = row.yes_step_id
                                no_raw = row.no_step_id

                                yes_val = None if pd.isna(yes_raw) else int(yes_raw)
                                no_val = None if pd.isna(no_raw) else int(no_raw)
                                for value, label in ((yes_val, "Yes"), (no_val, "No")):
                                    if value is not None and value not in valid_step_ids:
                                        raise ValueError(
                                            f"{label} → Step ID {value} does not exist yet. "
                                            "Use one of the Step ID values listed in the table."
                                        )
                                payload = {
                                    "step_order": int(order_raw),
                                    "title": title_raw.strip(),
                                    "yes_step_id": yes_val,
                                    "no_step_id": no_val,
                                }

                                if pd.isna(step_id):
                                    inserts.append(payload)
                                else:
                                    updates[int(step_id)] = payload

                            if to_delete:
                                delete_workflow_steps(to_delete)

                            if updates:
                                update_workflow_steps(updates)

                            workflow_id_int = int(workflow_id)
                            for payload in inserts:
                                insert_workflow_step(
                                    workflow_id_int,
                                    payload["step_order"],
                                    payload["title"],
                                    payload["yes_step_id"],
                                    payload["no_step_id"],
                                )

                            clear_cached_workflows()
                            st.success("Workflow table updated.")
                            st.session_state[state_key_source] = "table"
                            st.session_state[state_key_current] = None
                            trigger_rerun()
                        except ValueError as exc:
                            st.error(str(exc))

            with tab_barfi:
                render_barfi_editor(expander_label, int(workflow_id))


def render_add_topic() -> None:
    st.subheader("✍️ Log a new handover discussion")

    with st.form("add_topic_form", clear_on_submit=True):
        person = st.text_input("Person responsible", placeholder="e.g., Ahmed Hassan")
        topic = st.text_input("Topic", placeholder="e.g., Monthly demand planning handoff")
        meeting_date = st.date_input(
            "Target handover date", value=dt.date.today(), format="DD/MM/YYYY"
        )
        details = st.text_area(
            "Key details",
            placeholder="Summary, open questions, systems used, pending actions...",
        )
        status = st.selectbox(
            "Current status",
            STATUS_OPTIONS,
            index=0,
            format_func=status_badge,
        )
        next_steps = st.text_area(
            "Next steps / notes",
            placeholder="Follow-up tasks, documents to share, risks, etc.",
        )
        submitted = st.form_submit_button("Save topic", type="primary")
        if submitted:
            if not person.strip() or not topic.strip():
                st.error("Person and topic are required.")
            else:
                add_topic(person, topic, meeting_date, details, status, next_steps)
                clear_cached_topics()
                st.success("Topic saved and ready for review.")
                _trigger_rerun()


def _contact_select_options(df: pd.DataFrame) -> Dict[str, int]:
    return {
        f"#{int(row.id)} · {row.full_name} ({row.title or 'No title'})": int(row.id)
        for row in df.itertuples()
    }


def build_contacts_org_graph(
    all_contacts: pd.DataFrame,
    filtered: pd.DataFrame,
    font_size: int = 16,
) -> Optional[str]:
    if filtered.empty:
        return None

    required_columns = {"id", "full_name", "title", "manager_id"}
    if not required_columns.issubset(all_contacts.columns):
        return None

    def _safe_int(value: object) -> Optional[int]:
        try:
            if pd.isna(value):
                return None
        except TypeError:
            pass
        try:
            return int(value)
        except (TypeError, ValueError):
            return None

    id_to_row: Dict[int, pd.Series] = {}
    for row in all_contacts.itertuples(index=False):
        contact_id = _safe_int(getattr(row, "id", None))
        if contact_id is None:
            continue
        id_to_row[contact_id] = pd.Series(row._asdict())

    display_ids: Set[int] = set()
    queue: List[int] = []

    for row in filtered.itertuples(index=False):
        contact_id = _safe_int(getattr(row, "id", None))
        if contact_id is None:
            continue
        display_ids.add(contact_id)
        queue.append(contact_id)

    while queue:
        current_id = queue.pop()
        row = id_to_row.get(current_id)
        if row is None:
            continue
        manager_id = _safe_int(row.get("manager_id"))
        if manager_id is None or manager_id in display_ids:
            continue
        if manager_id in id_to_row:
            display_ids.add(manager_id)
            queue.append(manager_id)

    if not display_ids:
        return None

    parent_lookup: Dict[int, Optional[int]] = {}
    children_map: Dict[int, List[int]] = defaultdict(list)

    for contact_id in display_ids:
        row = id_to_row.get(contact_id)
        if row is None:
            continue
        manager_id = _safe_int(row.get("manager_id"))
        parent_lookup[contact_id] = manager_id if manager_id in display_ids else None
        if manager_id is not None and manager_id in display_ids:
            children_map[manager_id].append(contact_id)

    roots = [cid for cid in display_ids if parent_lookup.get(cid) is None]
    if not roots:
        roots = sorted(display_ids)

    levels: Dict[int, int] = {}
    bfs_queue: deque[Tuple[int, int]] = deque((root, 0) for root in roots)
    while bfs_queue:
        contact_id, level = bfs_queue.popleft()
        if contact_id in levels:
            continue
        levels[contact_id] = level
        for child_id in children_map.get(contact_id, []):
            bfs_queue.append((child_id, level + 1))

    graph_lines = [
        "digraph OrgStructure {",
        "    rankdir=TB;",
        '    graph [splines=ortho, nodesep=0.6, ranksep=0.9, margin="0.3,0.3"];',
        '    node [shape=box, style="rounded,filled", fontname="Helvetica", fontsize={}, fillcolor="#EDF2F7", color="#2D3748", fontcolor="#1A202C"];'.format(font_size),
        '    edge [color="#4A5568", arrowhead=normal];',
    ]

    filtered_ids = {int(idx) for idx in display_ids if idx in filtered.set_index("id").index}

    for contact_id in sorted(display_ids):
        row = id_to_row.get(contact_id)
        if row is None:
            continue
        name = str(row.get("full_name", "Unnamed"))
        title = str(row.get("title", "") or "")
        label_parts = [name.strip() or "Unnamed"]
        if title:
            label_parts.append(title.strip())
        label = "\n".join(label_parts).replace("\n\n", "\n")
        level = levels.get(contact_id, 0)
        fill = _org_level_color(level)
        border_color = "#6B46C1" if contact_id in filtered_ids else "#2D3748"
        penwidth = 3 if contact_id in filtered_ids else 1.8
        safe_label = label.replace("\"", "\\\"")
        graph_lines.append(
            f'    "node_{contact_id}" [label="{safe_label}", fillcolor="{fill}", color="{border_color}", penwidth={penwidth}];'
        )

    for contact_id in display_ids:
        row = id_to_row.get(contact_id)
        if row is None:
            continue
        manager_id = _safe_int(row.get("manager_id"))
        if manager_id is None or manager_id not in display_ids:
            continue
        graph_lines.append(f'    "node_{manager_id}" -> "node_{contact_id}";')

    graph_lines.append("}")
    return "\n".join(graph_lines)


def render_contacts_org_chart(all_contacts: pd.DataFrame, filtered: pd.DataFrame) -> None:
    graphviz_code = build_contacts_org_graph(all_contacts, filtered)
    if graphviz_code is None:
        st.info("Not enough hierarchy data to render the org chart.")
        return

    font_size = st.slider(
        "Org chart font size",
        min_value=10,
        max_value=26,
        value=16,
        help="Adjust label readability in the organization chart.",
    )
    adjusted_graph = build_contacts_org_graph(all_contacts, filtered, font_size=font_size)
    if adjusted_graph is None:
        st.info("No contacts available to chart after applying filters.")
        return

    st.graphviz_chart(adjusted_graph, use_container_width=True)


def render_contacts() -> None:
    st.subheader("📇 Contact book & org directory")

    contacts_df = load_contacts()

    manager_options = {"— None —": None}
    manager_options.update(_contact_select_options(contacts_df))

    with st.form("add_contact_form", clear_on_submit=True):
        full_name = st.text_input("Full name", placeholder="e.g., Sara Abdallah")
        col_title, col_department = st.columns(2)
        with col_title:
            title = st.text_input("Title", placeholder="e.g., BI Manager")
        with col_department:
            department = st.text_input("Department", placeholder="e.g., Business Intelligence")
        col_email, col_phone = st.columns(2)
        with col_email:
            email = st.text_input("Email", placeholder="name@example.com")
        with col_phone:
            phone = st.text_input("Phone", placeholder="+966 5 5555 5555")
        location = st.text_input("Location", placeholder="Headquarters / Remote")
        manager_label = st.selectbox("Manager", options=list(manager_options.keys()))
        manager_id = manager_options[manager_label]
        notes = st.text_area(
            "Notes",
            placeholder="Availability hours, preferred channels, responsibilities...",
        )

        submitted = st.form_submit_button("Add contact", type="primary")

    if submitted:
        if not full_name.strip():
            st.error("Full name is required.")
        else:
            add_contact(
                full_name,
                title,
                department,
                email,
                phone,
                location,
                manager_id,
                notes,
            )
            clear_cached_contacts()
            st.success("Contact added to the directory.")
            contacts_df = load_contacts()
            manager_options = {"— None —": None}
            manager_options.update(_contact_select_options(contacts_df))

    if contacts_df.empty:
        st.info("No contacts captured yet. Use the form above to add your first entry.")
        return

    with st.expander("Filters", expanded=False):
        departments = sorted(
            {
                str(dep).strip()
                for dep in contacts_df["department"].dropna()
                if str(dep).strip()
            }
        )
        locations = sorted(
            {
                str(loc).strip()
                for loc in contacts_df["location"].dropna()
                if str(loc).strip()
            }
        )
        selected_departments = st.multiselect("Departments", options=departments)
        selected_locations = st.multiselect("Locations", options=locations)
        show_with_manager = st.checkbox("Only contacts with a manager assigned", value=False)

    filtered = contacts_df.copy()
    if selected_departments:
        filtered = filtered[filtered["department"].isin(selected_departments)]
    if selected_locations:
        filtered = filtered[filtered["location"].isin(selected_locations)]
    if show_with_manager:
        filtered = filtered[filtered["manager_id"].notna()]

    st.markdown("### Directory")
    display_df = filtered.copy()
    manager_lookup = contacts_df.set_index("id")["full_name"].to_dict()
    display_df["manager_name"] = display_df["manager_id"].map(manager_lookup)
    display_df["created_at"] = display_df["created_at"].dt.strftime("%d %b %Y %H:%M")
    display_df["updated_at"] = display_df["updated_at"].dt.strftime("%d %b %Y %H:%M")
    selected_ids: List[int] = []
    st.caption("Toggle the checkbox to mark contacts for bulk removal before clicking delete.")
    table_payload = display_df[
        [
            "id",
            "full_name",
            "title",
            "department",
            "email",
            "phone",
            "location",
            "manager_name",
            "notes",
            "updated_at",
        ]
    ].copy()
    table_payload.insert(0, "Selected", False)
    edited_table = st.data_editor(
        table_payload,
        hide_index=True,
        use_container_width=True,
        disabled=[
            "id",
            "full_name",
            "title",
            "department",
            "email",
            "phone",
            "location",
            "manager_name",
            "notes",
            "updated_at",
        ],
        key="contacts_directory_editor",
    )
    if not edited_table.empty:
        selected_ids = edited_table.loc[edited_table["Selected"] == True, "id"].tolist()  # noqa: E712

    csv = display_df.to_csv(index=False)
    st.download_button(
        label="⬇️ Download contacts CSV",
        data=csv,
        file_name="contacts.csv",
        mime="text/csv",
        use_container_width=True,
    )

    st.markdown("### Organization structure")
    if filtered.empty:
        st.info("Add contacts (and assign managers) to display the org chart.")
    else:
        render_contacts_org_chart(contacts_df, filtered)

    st.markdown("### Edit or remove")
    options = _contact_select_options(contacts_df)
    if not options:
        st.info("No contacts to edit.")
        return

    selection = st.selectbox("Select a contact", options=list(options.keys()))
    contact_id = options[selection]
    target = contacts_df.loc[contacts_df["id"] == contact_id]
    if target.empty:
        st.warning("Contact not found; try refreshing.")
        return

    record = target.iloc[0]

    manager_picker_options = {"— None —": None}
    manager_picker_options.update(
        {
            label: cid
            for label, cid in _contact_select_options(contacts_df).items()
            if cid != contact_id
        }
    )
    default_manager_label = next(
        (label for label, cid in manager_picker_options.items() if cid == record.get("manager_id")),
        "— None —",
    )

    with st.form(f"update_contact_form_{contact_id}"):
        upd_full_name = st.text_input("Full name", value=record.get("full_name", ""))
        col_title, col_department = st.columns(2)
        with col_title:
            upd_title = st.text_input("Title", value=record.get("title", "") or "")
        with col_department:
            upd_department = st.text_input("Department", value=record.get("department", "") or "")
        col_email, col_phone = st.columns(2)
        with col_email:
            upd_email = st.text_input("Email", value=record.get("email", "") or "")
        with col_phone:
            upd_phone = st.text_input("Phone", value=record.get("phone", "") or "")
        upd_location = st.text_input("Location", value=record.get("location", "") or "")
        upd_manager_label = st.selectbox(
            "Manager",
            options=list(manager_picker_options.keys()),
            index=list(manager_picker_options.keys()).index(default_manager_label)
            if default_manager_label in manager_picker_options
            else 0,
        )
        upd_manager_id = manager_picker_options[upd_manager_label]
        upd_notes = st.text_area("Notes", value=record.get("notes", "") or "")

        save_changes = st.form_submit_button("Save changes", type="primary")

    if save_changes:
        if not upd_full_name.strip():
            st.error("Full name cannot be empty.")
        else:
            update_payload = {
                "full_name": upd_full_name.strip(),
                "title": upd_title.strip(),
                "department": upd_department.strip(),
                "email": upd_email.strip(),
                "phone": upd_phone.strip(),
                "location": upd_location.strip(),
                "manager_id": upd_manager_id,
                "notes": upd_notes.strip(),
            }
            update_contact(contact_id, update_payload)
            clear_cached_contacts()
            st.success("Contact updated.")
            _trigger_rerun()

    remove_disabled = not selected_ids
    if st.button(
        "🗑️ Remove selected contacts",
        type="secondary",
        disabled=remove_disabled,
        help="Select at least one row above before deleting.",
    ):
        delete_contacts(selected_ids)
        clear_cached_contacts()
        st.success("Selected contacts removed.")
        _trigger_rerun()


def render_backups() -> None:
    st.subheader("🗃️ Database backups")

    ensure_backup_directory()
    backups = list_database_backups()

    st.markdown(
        "Create dated snapshots of your data before big changes, or restore a previous copy if needed."
    )

    col_backup, col_refresh = st.columns([3, 1])
    with col_backup:
        if st.button("💾 Create backup", type="primary"):
            backup_path = create_database_backup()
            st.success(f"Backup created: {backup_path.name}")
            backups = list_database_backups()
    with col_refresh:
        if st.button("🔄 Refresh list"):
            backups = list_database_backups()
            st.info("Backup list refreshed.")

    if not backups:
        st.info("No backups found yet. Create your first snapshot to get started.")
        return

    options: List[Path] = []
    human_labels: List[str] = []
    for entry in backups:
        timestamp = _parse_backup_timestamp(entry)
        human = timestamp.strftime("%d %b %Y · %H:%M") if timestamp else entry.name
        options.append(entry)
        human_labels.append(f"{human} ({entry.name})")

    selected_index = st.selectbox(
        "Available backups",
        options=range(len(options)),
        format_func=lambda idx: human_labels[idx],
    )
    selected_backup = options[selected_index]

    st.caption(f"Selected file: {selected_backup}")

    st.divider()
    st.markdown("### Restore backup")
    confirm = st.checkbox(
        "I understand restoring will overwrite the current database with the selected backup.",
        value=False,
    )
    if st.button("⏪ Restore selected backup", type="primary", disabled=not confirm):
        try:
            restore_database_backup(selected_backup)
        except FileNotFoundError:
            st.error("Backup file no longer exists. Refresh the list and try again.")
        except Exception as exc:
            st.error(f"Failed to restore backup: {exc}")
        else:
            st.success("Backup restored. The app will refresh to show the restored data.")
            trigger_rerun()


def render_metrics(df: pd.DataFrame) -> None:
    total = len(df)
    in_progress = int((df["status"] == "In Progress").sum())
    blocked = int((df["status"] == "Waiting").sum())
    completed = int((df["status"] == "Completed").sum())

    upcoming_window = dt.date.today() + dt.timedelta(days=7)
    upcoming = int((df["meeting_date"] <= upcoming_window).sum())

    st.markdown("### Snapshot")
    col1, col2, col3, col4 = st.columns(4)
    col1.metric("🗂️ Total topics", total)
    col2.metric("🚧 In progress", in_progress)
    col3.metric("⏳ Waiting on others", blocked)
    col4.metric("⏰ Due ≤ 7 days", upcoming)

    completion_rate = completed / total if total else 0.0
    st.progress(
        completion_rate,
        text=f"Completion rate: {completed}/{total} done" if total else "Completion rate",
    )


def render_review() -> None:
    st.subheader("📋 Review & update handover topics")

    st.sidebar.markdown("### 💡 Quick tips")
    st.sidebar.success(
        "**Focus on clarity:** add document links, escalation paths, and who needs to be informed."
    )
    st.sidebar.info(
        "**Weekly rhythm:** update status, highlight blockers, and share the report link."
    )

    with st.expander("Filters", expanded=True):
        all_df = fetch_topics()
        persons = sorted(all_df["person"].unique()) if not all_df.empty else []
        statuses = sorted(all_df["status"].unique()) if not all_df.empty else STATUS_OPTIONS

        selected_persons = st.multiselect("People", options=persons)
        selected_statuses = st.multiselect(
            "Status",
            options=statuses,
            default=[],
            format_func=status_badge,
        )

    df = load_topics({"person": selected_persons, "status": selected_statuses})
    df_display = df

    if df.empty:
        st.info("No topics logged yet. Use the form above to add your first entry.")
        return

    render_metrics(df)

    if px is None:
        st.warning(
            "Plotly is not installed, so charts are hidden. Install `plotly` to see the visuals.",
            icon="⚠️",
        )
    else:
        status_summary = (
            df.groupby("status")["id"].count().reset_index().rename(columns={"id": "count"})
        )
        status_summary["display"] = status_summary["status"].apply(status_badge)

        pie_col, chart_col = st.columns([1, 2])
        with pie_col:
            fig = px.pie(
                status_summary,
                names="display",
                values="count",
                color="status",
                color_discrete_map=STATUS_COLORS,
                title="Status distribution",
            )
            st.plotly_chart(fig, use_container_width=True)

        with chart_col:
            timeline = (
                df.groupby("meeting_date")["id"].count().reset_index().rename(columns={"id": "Topics"})
            )
            fig_timeline = px.area(
                timeline,
                x="meeting_date",
                y="Topics",
                title="Upcoming handover timeline",
                color_discrete_sequence=["#636EFA"],
            )
            fig_timeline.update_traces(mode="lines+markers", line_shape="hv")
            st.plotly_chart(fig_timeline, use_container_width=True)

    st.markdown("### 📁 Handover register")
    st.dataframe(
        df_display.assign(
            meeting_date=df_display["meeting_date"].apply(lambda d: d.strftime("%d %b %Y")),
            created_at=df_display["created_at"].dt.strftime("%d %b %Y %H:%M"),
            updated_at=df_display["updated_at"].dt.strftime("%d %b %Y %H:%M"),
        ),
        use_container_width=True,
    )

    csv = df_display.to_csv(index=False)
    st.download_button(
        label="Download register (CSV)",
        data=csv,
        file_name="handover_topics.csv",
        mime="text/csv",
        use_container_width=True,
    )

    st.markdown("### 📤 Share & export")
    show_excel = EXCEL_ENGINE is not None
    columns = 3 if show_excel else 2
    layout = st.columns(columns)

    col_csv = layout[0]
    col_summary = layout[-1]
    excel_bytes = None

    if show_excel:
        col_excel = layout[1]
        try:
            excel_bytes = build_excel_report(df)
        except RuntimeError:
            show_excel = False

    summary_markdown = build_summary(df)

    with col_csv:
        st.download_button(
            label="⬇️ CSV",
            data=csv,
            file_name="handover_topics.csv",
            mime="text/csv",
            use_container_width=True,
        )

    if show_excel and excel_bytes:
        with col_excel:  # type: ignore[name-defined]
            st.download_button(
                label="📊 Excel report",
                data=excel_bytes,
                file_name="handover_report.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
    elif not show_excel:
        st.info(
            "Install `xlsxwriter` or `openpyxl` to enable Excel exports.",
            icon="ℹ️",
        )

    with col_summary:
        st.download_button(
            label="📝 Markdown summary",
            data=summary_markdown,
            file_name="handover_summary.md",
            mime="text/markdown",
            use_container_width=True,
        )

    share_subject = quote("Handover status update")
    share_body = quote(summary_markdown)
    st.markdown(
        f"[📧 Share summary via email](mailto:?subject={share_subject}&body={share_body})",
        help="Opens your email client with the summary pre-filled.",
    )

    st.markdown("### Update an entry")
    topic_options = {
        f"#{row.id} · {row.topic} ({row.person})": row.id for row in df.itertuples()
    }

    if not topic_options:
        st.info("No topics to update.")
        return

    selection = st.selectbox("Select a topic", options=list(topic_options.keys()))
    topic_id = topic_options[selection]
    selected_topic = fetch_topic(topic_id)

    if selected_topic is None:
        st.warning("Topic not found; it may have been removed.")
        return

    with st.form("update_topic_form"):
        new_person = st.text_input("Person responsible", value=selected_topic.person)
        new_topic = st.text_input("Topic", value=selected_topic.topic)
        new_meeting_date = st.date_input(
            "Target handover date",
            value=selected_topic.meeting_date,
            format="DD/MM/YYYY",
        )
        new_status = st.selectbox(
            "Status",
            STATUS_OPTIONS,
            index=STATUS_OPTIONS.index(selected_topic.status)
            if selected_topic.status in STATUS_OPTIONS
            else 0,
            format_func=status_badge,
        )
        new_details = st.text_area("Key details", value=selected_topic.details)
        new_next_steps = st.text_area("Next steps / notes", value=selected_topic.next_steps)

        submitted = st.form_submit_button("Apply changes", type="primary")

    if submitted:
        updates = {
            "person": new_person.strip(),
            "topic": new_topic.strip(),
            "details": new_details.strip(),
            "status": new_status,
            "next_steps": new_next_steps.strip(),
        }
        update_topic(topic_id, updates, meeting_date=new_meeting_date)
        clear_cached_topics()
        st.success("Topic updated.")


def render_stakeholders() -> None:
    st.subheader("🤝 Stakeholder analysis & management")

    contacts_df = load_contacts()
    contact_choices: Dict[str, Optional[int]] = {"— Manual entry —": None}
    if not contacts_df.empty:
        for row in contacts_df.itertuples():
            label = f"{row.full_name} ({row.title or 'No title'})"
            contact_choices[label] = int(row.id)

    with st.form("add_stakeholder_form", clear_on_submit=True):
        selected_contact_label = st.selectbox(
            "Select from contacts",
            options=list(contact_choices.keys()),
            help="Pick an existing contact to auto-fill name/role, or choose manual entry.",
        )
        linked_contact_id = contact_choices[selected_contact_label]

        if linked_contact_id is not None and not contacts_df.empty:
            match = contacts_df.loc[contacts_df["id"] == linked_contact_id].iloc[0]
            prefill_name = match.get("full_name", "")
            prefill_position = match.get("title", "") or ""
            prefill_department = match.get("department", "") or ""
        else:
            prefill_name = ""
            prefill_position = ""
            prefill_department = ""
        name = st.text_input("Name", value=prefill_name, placeholder="e.g., Chief Operations Officer")
        position = st.text_input("Role / position", value=prefill_position, placeholder="e.g., COO")
        department = st.text_input("Department / team", value=prefill_department, placeholder="e.g., Operations")
        role_description = st.text_area(
            "Engagement role",
            placeholder="Decision-maker, champion, subject-matter expert, etc.",
        )
        influence = st.slider(
            "Influence (1 = low · 5 = high)",
            min_value=1,
            max_value=5,
            value=4,
        )
        interest = st.slider(
            "Interest (1 = low · 5 = high)",
            min_value=1,
            max_value=5,
            value=4,
        )
        power = st.slider(
            "Power (1 = low · 5 = high)",
            min_value=1,
            max_value=5,
            value=4,
        )
        support_level = st.selectbox(
            "Support level",
            STAKEHOLDER_SUPPORT_LEVELS,
            index=0,
        )
        notes = st.text_area(
            "Notes",
            placeholder="Key expectations, engagement tactics, meeting cadence...",
        )

        submitted = st.form_submit_button("Add stakeholder", type="primary")

    if submitted:
        if not name.strip():
            st.error("Name is required.")
        else:
            add_stakeholder(
                name,
                position,
                department,
                role_description,
                influence,
                interest,
                power,
                support_level,
                notes,
                contact_id=linked_contact_id,
            )
            clear_cached_stakeholders()
            st.success("Stakeholder added to the register.")

    df = load_stakeholders()

    if df.empty:
        st.info("No stakeholders logged yet. Use the form above to add your first entry.")
        return

    with st.expander("Filters", expanded=False):
        departments = sorted(
            {
                str(dep).strip()
                for dep in df["department"].dropna()
                if str(dep).strip()
            }
        )
        support_levels = STAKEHOLDER_SUPPORT_LEVELS
        selected_departments = st.multiselect("Departments", options=departments)
        selected_support = st.multiselect(
            "Support level",
            options=support_levels,
            default=[],
        )
        min_power, min_interest = st.columns(2)
        with min_power:
            min_power_value = st.slider(
                "Minimum power",
                min_value=1,
                max_value=5,
                value=1,
            )
        with min_interest:
            min_interest_value = st.slider(
                "Minimum interest",
                min_value=1,
                max_value=5,
                value=1,
            )

    filtered = df.copy()
    if selected_departments:
        filtered = filtered[filtered["department"].isin(selected_departments)]
    if selected_support:
        filtered = filtered[filtered["support_level"].isin(selected_support)]
    filtered = filtered[
        (filtered["power"].fillna(0) >= min_power_value)
        & (filtered["interest"].fillna(0) >= min_interest_value)
    ]

    if filtered.empty:
        st.warning("No stakeholders match the current filters.")
        return

    col_total, col_advocates, col_resistant, col_high_power = st.columns(4)
    col_total.metric("Total stakeholders", len(filtered))
    col_advocates.metric(
        "Advocates",
        int((filtered["support_level"] == "Advocate").sum()),
    )
    col_resistant.metric(
        "Resistant",
        int((filtered["support_level"] == "Resistant").sum()),
    )
    col_high_power.metric(
        "High-power (≥4)",
        int((filtered["power"] >= 4).sum()),
    )

    if px is None:
        st.info(
            "Install `plotly` to see the power vs. interest bubble chart.",
            icon="ℹ️",
        )
    else:
        scatter_df = filtered.copy()
        scatter_df["interest"] = scatter_df["interest"].astype(float)
        scatter_df["power"] = scatter_df["power"].astype(float)
        scatter_df["influence"] = scatter_df["influence"].astype(float)
        fig = px.scatter(
            scatter_df,
            x="interest",
            y="power",
            size="influence",
            color="support_level",
            hover_data={
                "name": True,
                "position": True,
                "department": True,
                "role_description": True,
                "support_level": True,
                "interest": True,
                "power": True,
            },
            size_max=40,
            color_discrete_map={
                "Advocate": "#38A169",
                "Neutral": "#ECC94B",
                "Resistant": "#E53E3E",
            },
            title="Power vs. interest map",
        )
        fig.update_layout(
            xaxis=dict(title="Interest", range=[0.5, 5.5], dtick=1),
            yaxis=dict(title="Power", range=[0.5, 5.5], dtick=1),
            legend_title_text="Support level",
        )
        st.plotly_chart(fig, use_container_width=True)

    display_df = filtered.copy()
    if "contact_id" in display_df.columns and not contacts_df.empty:
        contact_name_lookup = contacts_df.set_index("id")["full_name"].to_dict()
        display_df["contact_name"] = display_df["contact_id"].map(contact_name_lookup)
    display_df["created_at"] = display_df["created_at"].dt.strftime("%d %b %Y %H:%M")
    display_df["updated_at"] = display_df["updated_at"].dt.strftime("%d %b %Y %H:%M")

    st.markdown("### Stakeholder register")
    st.dataframe(display_df, use_container_width=True)

    csv = display_df.to_csv(index=False)
    st.download_button(
        label="⬇️ Download CSV",
        data=csv,
        file_name="stakeholders.csv",
        mime="text/csv",
        use_container_width=True,
    )

    options = {
        f"#{row.id} · {row.name} ({row.position or 'No title'})": int(row.id)
        for row in df.itertuples()
    }

    st.markdown("### Update or remove")
    if not options:
        st.info("No stakeholders available to edit.")
        return

    selection = st.selectbox("Select a stakeholder", options=list(options.keys()))
    stakeholder_id = options[selection]
    target = df.loc[df["id"] == stakeholder_id]
    if target.empty:
        st.warning("Stakeholder not found; try refreshing.")
        return

    record = target.iloc[0]

    contact_update_choices: Dict[str, Optional[int]] = {"— Manual entry —": None}
    if not contacts_df.empty:
        for row in contacts_df.itertuples():
            label = f"{row.full_name} ({row.title or 'No title'})"
            contact_update_choices[label] = int(row.id)
    current_contact_raw = record.get("contact_id")
    current_contact_id: Optional[int]
    if pd.isna(current_contact_raw):
        current_contact_id = None
    else:
        try:
            current_contact_id = int(current_contact_raw)
        except (TypeError, ValueError):
            current_contact_id = None
    default_label = next(
        (label for label, cid in contact_update_choices.items() if cid == current_contact_id),
        "— Manual entry —",
    )

    with st.form(f"update_stakeholder_form_{stakeholder_id}"):
        upd_name = st.text_input("Name", value=str(record["name"]))
        upd_position = st.text_input("Role / position", value=record.get("position", "") or "")
        upd_department = st.text_input("Department / team", value=record.get("department", "") or "")
        upd_role_desc = st.text_area(
            "Engagement role",
            value=record.get("role_description", "") or "",
        )
        upd_influence = st.slider(
            "Influence",
            min_value=1,
            max_value=5,
            value=int(record.get("influence", 3) or 3),
        )
        upd_interest = st.slider(
            "Interest",
            min_value=1,
            max_value=5,
            value=int(record.get("interest", 3) or 3),
        )
        upd_power = st.slider(
            "Power",
            min_value=1,
            max_value=5,
            value=int(record.get("power", 3) or 3),
        )
        current_support = record.get("support_level", STAKEHOLDER_SUPPORT_LEVELS[0])
        support_index = (
            STAKEHOLDER_SUPPORT_LEVELS.index(current_support)
            if current_support in STAKEHOLDER_SUPPORT_LEVELS
            else 0
        )
        upd_support = st.selectbox(
            "Support level",
            STAKEHOLDER_SUPPORT_LEVELS,
            index=support_index,
        )
        upd_notes = st.text_area("Notes", value=record.get("notes", "") or "")
        upd_contact_label = st.selectbox(
            "Linked contact",
            options=list(contact_update_choices.keys()),
            index=list(contact_update_choices.keys()).index(default_label)
            if default_label in contact_update_choices
            else 0,
        )
        upd_contact_id = contact_update_choices[upd_contact_label]

        save_changes = st.form_submit_button("Save changes", type="primary")

    if save_changes:
        if not upd_name.strip():
            st.error("Name cannot be empty.")
        else:
            update_stakeholder(
                stakeholder_id,
                {
                    "name": upd_name.strip(),
                    "position": upd_position.strip(),
                    "department": upd_department.strip(),
                    "role_description": upd_role_desc.strip(),
                    "influence": int(upd_influence),
                    "interest": int(upd_interest),
                    "power": int(upd_power),
                    "support_level": upd_support,
                    "notes": upd_notes.strip(),
                    "contact_id": upd_contact_id,
                },
            )
            clear_cached_stakeholders()
            st.success("Stakeholder updated.")

    if st.button("🗑️ Delete selected stakeholder", type="secondary"):
        delete_stakeholders([stakeholder_id])
        clear_cached_stakeholders()
        st.success("Stakeholder removed.")
        st.experimental_rerun()


def main() -> None:
    st.set_page_config(page_title="Handover Tracker", layout="wide")

    ship_positions = {
        "Contacts": -15,
        "Add topic": 10,
        "Review & update": 35,
        "Workflow builder": 65,
        "Stakeholders": 95,
        "Backups": 120,
    }
    initial_ship_pos = ship_positions["Contacts"]
    st.markdown(
        f"""
        <style>
            :root {{
                font-size: 18px;
                --ship-pos: {initial_ship_pos}%;
            }}
            body, .block-container {{
                font-size: 1.1rem !important;
            }}
            h1, h2, h3, h4, h5, h6 {{
                font-weight: 600;
                line-height: 1.35;
            }}
            .stMarkdown p,
            .stMarkdown li,
            div[data-testid="stMetricValue"],
            div[data-testid="stMetricLabel"],
            .stDataFrame,
            label,
            .stTextInput input,
            .stTextArea textarea,
            div[data-baseweb="select"],
            button,
            .stButton button {{
                font-size: 1.15rem !important;
            }}
            .ship-banner {{
                position: relative;
                width: 100%;
                height: 80px;
                overflow: hidden;
                margin-bottom: 0.75rem;
            }}
            .ship-banner .ship-icon {{
                position: absolute;
                top: 8px;
                left: 0;
                transform: translateX(var(--ship-pos));
                transition: transform 1.8s cubic-bezier(0.45, 0, 0.55, 1);
                width: 140px;
            }}
            .ship-banner .island {{
                position: absolute;
                bottom: -4px;
                right: 6%;
                width: 170px;
                animation: islandSway 5.5s ease-in-out infinite;
            }}
            @keyframes islandSway {{
                0% {{
                    transform: translateY(0);
                }}
                50% {{
                    transform: translateY(-3px);
                }}
                100% {{
                    transform: translateY(0);
                }}
            }}
        </style>
        """,
        unsafe_allow_html=True,
    )
    st.markdown(
        """
        <div class="ship-banner">
            <div class="ship-icon">
                <svg width="140" height="60" viewBox="0 0 140 60" xmlns="http://www.w3.org/2000/svg">
                    <defs>
                        <linearGradient id="hullGradient" x1="0%" y1="0%" x2="100%" y2="0%">
                            <stop offset="0%" stop-color="#2C5282" />
                            <stop offset="100%" stop-color="#63B3ED" />
                        </linearGradient>
                        <linearGradient id="sailGradient" x1="0%" y1="0%" x2="0%" y2="100%">
                            <stop offset="0%" stop-color="#FFFFFF" stop-opacity="0.95" />
                            <stop offset="100%" stop-color="#E2E8F0" stop-opacity="0.9" />
                        </linearGradient>
                    </defs>
                    <rect x="10" y="40" width="120" height="12" rx="6" fill="url(#hullGradient)" />
                    <polygon points="20,40 50,15 65,40" fill="url(#sailGradient)" stroke="#A0AEC0" stroke-width="1" />
                    <polygon points="65,40 95,20 95,40" fill="url(#sailGradient)" stroke="#A0AEC0" stroke-width="1" />
                    <rect x="63" y="18" width="4" height="24" fill="#4A5568" />
                    <circle cx="110" cy="30" r="6" fill="#F6E05E" opacity="0.85">
                        <animate attributeName="opacity" values="0.6;0.95;0.6" dur="4s" repeatCount="indefinite" />
                    </circle>
                    <path d="M15 50 Q35 45 55 50 T95 50 T135 50" stroke="#90CDF4" stroke-width="3" fill="none" stroke-linecap="round">
                        <animate attributeName="d" dur="6s" repeatCount="indefinite"
                            values="M15 50 Q35 45 55 50 T95 50 T135 50;
                                   M15 50 Q35 47 55 50 T95 53 T135 50;
                                   M15 50 Q35 45 55 50 T95 50 T135 50" />
                    </path>
                </svg>
            </div>
            <div class="island">
                <svg width="170" height="80" viewBox="0 0 170 80" xmlns="http://www.w3.org/2000/svg">
                    <defs>
                        <linearGradient id="islandSand" x1="0%" y1="0%" x2="0%" y2="100%">
                            <stop offset="0%" stop-color="#FCD9A5" />
                            <stop offset="100%" stop-color="#F6AD55" />
                        </linearGradient>
                        <linearGradient id="islandGrass" x1="0%" y1="0%" x2="0%" y2="100%">
                            <stop offset="0%" stop-color="#48BB78" />
                            <stop offset="100%" stop-color="#2F855A" />
                        </linearGradient>
                    </defs>
                    <ellipse cx="85" cy="62" rx="75" ry="14" fill="#63B3ED" opacity="0.45" />
                    <path d="M25 55 C55 35 115 35 145 55" fill="url(#islandSand)" />
                    <path d="M35 53 C60 38 110 38 135 53" fill="url(#islandGrass)" />
                    <g>
                        <rect x="78" y="12" width="8" height="32" fill="#744210" rx="2" />
                        <path d="M82 10 C68 4 60 2 48 7 C65 12 70 17 72 24" fill="#48BB78" />
                        <path d="M82 10 C96 2 105 0 118 6 C102 11 97 17 95 24" fill="#38A169" />
                        <circle cx="54" cy="42" r="5" fill="#68D391" />
                        <circle cx="110" cy="40" r="6" fill="#34D399" />
                    </g>
                    <g>
                        <circle cx="64" cy="50" r="3.2" fill="#F687B3" />
                        <circle cx="68" cy="48" r="2.4" fill="#ED64A6" />
                        <circle cx="118" cy="49" r="3" fill="#F6ADCD" />
                        <circle cx="122" cy="47" r="2.2" fill="#ED64A6" />
                    </g>
                </svg>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    st.title("Onboarding command deck")
    st.caption("Chart your contacts, handovers, and workflows for a smooth landing.")

    init_db()

    (
        tab_contacts,
        tab_add,
        tab_review,
        tab_workflow,
        tab_stakeholders,
        tab_backups,
    ) = st.tabs(
        [
            "Contacts",
            "Add topic",
            "Review & update",
            "Workflow builder",
            "Stakeholders",
            "Backups",
        ]
    )

    ship_positions_js = json.dumps(ship_positions)
    st.markdown(
        f"""
        <script>
        const shipPositions = {ship_positions_js};
        function updateShipPosition() {{
            const externalDoc = window.parent?.document ?? document;
            const tabs = externalDoc.querySelectorAll('button[data-baseweb="tab"]');
            tabs.forEach(tab => {{
                if (tab.getAttribute('aria-selected') === 'true') {{
                    const label = tab.textContent.trim();
                    const offset = shipPositions[label] ?? -15;
                    externalDoc.documentElement.style.setProperty('--ship-pos', offset + '%');
                }}
            }});
        }}
        const externalDoc = window.parent?.document ?? document;
        if (externalDoc.readyState === 'loading') {{
            externalDoc.addEventListener('DOMContentLoaded', () => setTimeout(updateShipPosition, 150));
        }} else {{
            setTimeout(updateShipPosition, 150);
        }}
        externalDoc.addEventListener('click', evt => {{
            if (evt.target.closest('button[data-baseweb="tab"]')) {{
                setTimeout(updateShipPosition, 160);
            }}
        }});
        setInterval(updateShipPosition, 2000);
        </script>
        """,
        unsafe_allow_html=True,
    )

    with tab_contacts:
        render_contacts()

    with tab_add:
        render_add_topic()

    with tab_review:
        render_review()

    with tab_workflow:
        render_workflows()

    with tab_stakeholders:
        password = st.text_input("Enter password to view Stakeholders", type="password")
        if password == "20sara29":
            render_stakeholders()
        else:
            st.markdown(":)")

    with tab_backups:
        render_backups()


if __name__ == "__main__":
    main()
